"""Bounded, in-memory public procurement document extraction with evidence locators."""
import io
import ipaddress
import re
import zipfile
from pathlib import PurePosixPath
from urllib.parse import urljoin, urlparse, unquote

MAX_BYTES = 30 * 1024 * 1024
MAX_MEMBERS = 100
MAX_EXPANDED = 60 * 1024 * 1024
ATTACHMENTS = ('.pdf', '.doc', '.docx', '.xls', '.xlsx', '.zip', '.rar', '.txt')
GENERIC_TITLE = re.compile(r'^(?:公告内容文档|公告|采购公告|招标公告|竞价公告|采购需求|采购意向|.*政府采购网)$')


def _block(out, text, locator, kind='paragraph'):
    text = re.sub(r'[\t\r ]+', ' ', str(text)).strip()
    if text:
        out['blocks'].append({'text': text, 'locator': locator, 'kind': kind})


def _public_link(target):
    try:
        parsed=urlparse(target)
        host=(parsed.hostname or '').rstrip('.').lower()
        if parsed.scheme not in ('http','https') or not host or parsed.username is not None or parsed.password is not None:
            return False
        if host=='localhost' or host.endswith(('.localhost','.local','.internal')):
            return False
        try: return ipaddress.ip_address(host).is_global
        except ValueError: return not re.fullmatch(r'[\d.]+',host)
    except ValueError:
        return False


def _project_title(out):
    """Use an explicitly labeled project/contract name, never a contact section."""
    headers={}; candidates=[]
    for block in out['blocks']:
        text=block['text']
        match=re.search(r'(?:采购项目名称|项目名称|合同名称)\s*[:：|]\s*([^\n|]{4,180})',text)
        if match: candidates.append((match.group(1).strip(),block))
        if block['kind']=='table_row':
            key=block['locator'].rsplit('/row:',1)[0]
            cells=[c.strip() for c in text.split('|')]
            named=[i for i,c in enumerate(cells) if c in ('采购项目名称','项目名称','合同名称')]
            if named: headers[key]=named
            elif key in headers:
                candidates.extend((cells[i],block) for i in headers[key] if i<len(cells) and len(cells[i])>=4)
    names={name for name,_ in candidates if not GENERIC_TITLE.fullmatch(name) and name not in ('采购需求概况','采购单位名称','采购项目名称','预算金额','预计采购时间')}
    if len(names)==1:
        name=names.pop(); block=next(b for value,b in candidates if value==name)
        out['title']=name
        out['title_evidence']={'text':block['text'],'locator':block['locator'],'method':'explicit_project_name'}


def _html(data, url, out):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(data, 'html.parser', from_encoding='utf-8' if _utf8(data) else None)
    for node in soup.select('script, style, nav, footer, header, aside, noscript, .nav, .navigation, .footer, .breadcrumb, .new_miew, .fixedbox'):
        node.decompose()
    # A selector union returns the earliest DOM match, often a search widget.
    # Prefer explicit notice containers, then generic layout classes.
    selectors=('#zbDiv','article','.article','main','#zoom','#content','.article-content','.vF_detail_content','.vT_detail_content','.TRS_Editor','.detail-content','.content')
    root=None; empty_root=None
    for selector in selectors:
        candidates=soup.select(selector)
        for node in candidates:
            value=node.get_text(' ',strip=True)
            if not value and selector!='.content': empty_root=empty_root if empty_root is not None else node
            if len(value)<100 and re.search(r'立即登录|登录/注册|开启全网商机',value):continue
            if value and not re.fullmatch(r'[\s/]*(?:(?:搜标题|搜全文|搜索|查询)[\s/]*)+',value):
                root=node; break
        if root is not None: break
    structured=root is not None
    if root is None: root=empty_root if empty_root is not None else (soup.body or soup)
    headings=list(root.find_all(['h1','h2']))+list(soup.select('h1, h2.tc, .info-title'))
    titles=[]
    for node in headings+([soup.title] if soup.title else []):
        value=node.get_text(' ',strip=True)
        if not value or GENERIC_TITLE.fullmatch(value): continue
        if node.name!='title' and (re.match(r'^(?:\d+|[一二三四五六七八九十]+)[.、．]\s*',value) or re.fullmatch(r'(?:项目联系方式|采购人信息|采购代理机构信息|联系方式|项目基本情况)',value)):
            continue
        titles.append(((bool(re.search(r'项目|采购|招标|磋商|中标|成交',value)),node.name!='title',len(value)),value,node))
    if titles:
        _,out['title'],title_node=max(titles,key=lambda item:item[0])
        out['title_evidence']={'text':out['title'],'locator':'html/'+title_node.name,'method':'heading'}
    else:
        out['title']=soup.title.get_text(' ',strip=True) if soup.title else ''
    for i,node in enumerate(soup.find_all(string=re.compile(r'发布时间|发布日期|公告日期|信息时间')),1):
        if root not in node.parents and len(node.strip())<250:
            value=str(node).strip()
            for parent in node.parents:
                if parent is root or parent.name in ('html','body','[document]') or root in parent.descendants:
                    break
                candidate=parent.get_text(' ',strip=True)
                if len(candidate)>250: break
                value=candidate
                if re.search(r'\d{4}[年./-]\d{1,2}[月./-]\d{1,2}',candidate): break
            _block(out,value,f'html/metadata:{i}','metadata')
    for i,node in enumerate(soup.select('.conttime, .time, .titx, .publish-time, time'),1):
        value=node.get_text(' ',strip=True)
        if node.name=='time' and node.get('datetime'): value='发布时间：'+node['datetime']
        if len(value)<=250 and re.search(r'\d{4}[年./-]\d{1,2}[月./-]\d{1,2}',value) and not re.search(r'服务|履约|部署|截止|开标|开启|响应|报名',value):
            _block(out,value,f'html/time-widget:{i}','metadata')
    for i,node in enumerate(soup.find_all('meta'),1):
        if node.get('property',node.get('name','')) in ('article:published_time','publishdate','pubdate') and node.get('content'):
            _block(out,'发布时间：'+node['content'],f'html/meta:{i}','metadata')
    for a in soup.find_all('a', href=True):
        target = urljoin(url, a['href'])
        if not _public_link(target): continue
        label = a.get_text(' ', strip=True)
        kind = 'attachment' if urlparse(target).path.lower().endswith(ATTACHMENTS) or re.search('附件|下载|采购文件|招标文件', label) else 'page'
        if root not in a.parents and kind!='attachment': continue
        out['links'].append({'url':target, 'text':label, 'kind':kind})
    frames=list(root.find_all('iframe'))
    for frame in frames:
        target=urljoin(url,frame.get('src',''))
        if frame.get('src') and _public_link(target):
            kind='attachment' if urlparse(target).path.lower().endswith(ATTACHMENTS) else 'page'
            out['links'].append({'url':target,'text':frame.get('title') or 'iframe 公告正文入口','kind':kind})
    if frames: out['warnings'].append('iframe content not fetched; public entry links retained for follow-up')
    body_text=' '.join(str(n).strip() for n in root.find_all(string=True) if n.parent.name not in ('h1','h2','title') and not any(p.name in ('a','iframe') for p in n.parents))
    identity=re.search(r'(?:采购项目名称|项目名称|合同名称|项目编号|采购人(?:名称)?)\s*[:：|]',body_text)
    unavailable=(root is empty_root or (frames and not body_text.strip()) or (not structured and GENERIC_TITLE.fullmatch(out['title'] or '') and not identity))
    out['content_status']='unavailable' if unavailable else 'available'
    if unavailable:
        out['status']='partial'
        out['warnings'].append('notice body unavailable: empty, dynamic, or listing page; retain search lead and public links')
        if out['title']: _block(out,out['title'],'html/notice-title','metadata')
    tables = list(root.find_all('table'))
    # Extract child tables before removing their layout-table ancestors.
    for ti, table in sorted(enumerate(tables,1),key=lambda item:len(list(item[1].parents)),reverse=True):
        grid = {}
        for ri, row in enumerate(table.find_all('tr'), 1):
            if row.find_parent('table') is not table: continue
            ci = 1
            for cell in row.find_all(['td','th'], recursive=False):
                while (ri,ci) in grid: ci += 1
                text = cell.get_text(' ',strip=True)
                try: rs,cs = min(int(cell.get('rowspan',1)),100),min(int(cell.get('colspan',1)),100)
                except ValueError: rs,cs=1,1
                for r in range(ri,ri+max(rs,1)):
                    for c in range(ci,ci+max(cs,1)): grid[r,c]=text
                ci += max(cs,1)
            vals=[grid[r,c] for r,c in sorted(grid) if r==ri]
            _block(out,' | '.join(vals),f'html/table:{ti}/row:{ri}','table_row')
        table.decompose()
    for i, line in enumerate(root.get_text('\n',strip=True).splitlines(),1):
        _block(out,line,f'html/line:{i}')
    if not out['title'] or GENERIC_TITLE.fullmatch(out['title']):
        _project_title(out)
        if out.get('title_evidence',{}).get('method')=='explicit_project_name':
            out['content_status']='available'
            if unavailable:
                out['status']='ok'
                out['warnings']=[w for w in out['warnings'] if not w.startswith('notice body unavailable:')]


    retained='\n'.join(b['text'] for b in out['blocks'])
    gated=re.search(r'查看(?:政府采购)?详细信息[，,\s]*(?:注册|请|登录)|本网站会员请[\s\S]{0,80}(?:登录|注册)|仅对.{0,20}会员开放|下文中.{0,30}为隐藏内容|登录后(?:查看|可见)',retained)
    masked=bool(re.search(r'\*{3,}|(?:公司|项目|采购人)\s*[（(]略[）)]',retained))
    if gated and (len(retained)<1500 or masked):
        out['content_status']='unavailable';out['status']='partial'
        out['warnings'].append('member-only or redacted summary; detailed notice body not acquired')


def _utf8(data):
    try: data.decode('utf-8'); return True
    except UnicodeDecodeError: return False


def _zip_safe(data):
    z=zipfile.ZipFile(io.BytesIO(data)); infos=z.infolist()
    if len(infos)>MAX_MEMBERS or sum(i.file_size for i in infos)>MAX_EXPANDED:
        z.close(); raise ValueError('archive expansion/member limit exceeded')
    return z


def _parse(data,url,content_type,depth):
    out={'title':'','text':'','blocks':[],'links':[],'status':'ok','warnings':[]}
    if len(data)>MAX_BYTES:
        out.update(status='size_limit',warnings=['document exceeds size limit']); return out
    ext=PurePosixPath(unquote(urlparse(url).path)).suffix.lower()
    ctype=content_type.lower()
    try:
        if 'wordprocessingml' in ctype: ext='.docx'
        elif 'spreadsheetml' in ctype: ext='.xlsx'
        elif data.startswith(b'PK'):
            with _zip_safe(data) as archive:
                names=set(archive.namelist())
                if 'word/document.xml' in names: ext='.docx'
                elif 'xl/workbook.xml' in names: ext='.xlsx'
                else: ext='.zip'
        if ext=='.pdf' or 'application/pdf' in ctype or data.startswith(b'%PDF'):
            try:
                from pypdf import PdfReader
                reader=PdfReader(io.BytesIO(data))
                texts=[page.extract_text() or '' for page in reader.pages]
            except ImportError:
                import pdfplumber
                with pdfplumber.open(io.BytesIO(data)) as reader:
                    texts=[page.extract_text() or '' for page in reader.pages]
            for i,text in enumerate(texts,1): _block(out,text,f'pdf/page:{i}','page')
            if not texts or any(not t.strip() for t in texts):
                out.update(status='needs_ocr',warnings=['PDF contains pages without extractable text; OCR required'])
        elif ext in ('.docx','.xlsx'):
            with _zip_safe(data): pass
            if ext=='.docx':
                from docx import Document
                doc=Document(io.BytesIO(data))
                for i,p in enumerate(doc.paragraphs,1): _block(out,p.text,f'docx/paragraph:{i}')
                for ti,t in enumerate(doc.tables,1):
                    for ri,row in enumerate(t.rows,1): _block(out,' | '.join(c.text for c in row.cells),f'docx/table:{ti}/row:{ri}','table_row')
            else:
                from openpyxl import load_workbook
                book=load_workbook(io.BytesIO(data),data_only=True)
                for sheet in book:
                    if sheet.max_row * sheet.max_column > 200000: raise ValueError('worksheet cell limit exceeded')
                    merged={}
                    for area in sheet.merged_cells.ranges:
                        for row in sheet.iter_rows(min_row=area.min_row,max_row=area.max_row,min_col=area.min_col,max_col=area.max_col):
                            for cell in row: merged[cell.coordinate]=sheet.cell(area.min_row,area.min_col).value
                    for ri,row in enumerate(sheet.iter_rows(),1):
                        vals=[merged.get(c.coordinate,c.value) for c in row]
                        if any(v is not None for v in vals): _block(out,' | '.join('' if v is None else str(v) for v in vals),f'xlsx/sheet:{sheet.title}/row:{ri}','table_row')
                book.close()
        elif ext=='.zip' or 'application/zip' in ctype:
            if depth>=2: raise ValueError('nested archive limit exceeded')
            with _zip_safe(data) as z:
                for info in z.infolist():
                    if info.is_dir(): continue
                    name=info.filename
                    if name.startswith(('/', '\\')) or '..' in PurePosixPath(name.replace('\\','/')).parts or re.match(r'^[A-Za-z]:',name):
                        out['warnings'].append(f'unsafe member skipped: {name}'); out['status']='partial'; continue
                    child=_parse(z.read(info),name,'',depth+1)
                    for b in child['blocks']: out['blocks'].append(dict(b,locator=f'zip/{name}/{b["locator"]}'))
                    if child['status']!='ok':
                        out['status']='partial'; out['warnings'].append(f'{name}: {child["status"]}')
                    out['warnings'].extend(f'{name}: {w}' for w in child['warnings'])
        elif ext in ('.doc','.xls','.rar'):
            out.update(status='unsupported',warnings=[f'legacy format {ext} requires external conversion'])
        elif 'html' in ctype or ext in ('.html','.htm','.shtml') or re.search(br'<(?:!doctype|html|body|article|div|p)[\s>]',data[:4096],re.I):
            _html(data,url,out)
        elif ext in ('.txt','.csv','.md') or ctype.startswith('text/'):
            try: text=data.decode('utf-8-sig')
            except UnicodeDecodeError: text=data.decode('gb18030')
            for i,line in enumerate(text.splitlines(),1): _block(out,line,f'text/line:{i}')
        else:
            out.update(status='unsupported',warnings=['unrecognized document format'])
    except ImportError as exc:
        out.update(status='dependency_missing'); out['warnings'].append(f'optional dependency unavailable: {exc.name}')
    except Exception as exc:
        out.update(status='parse_error'); out['warnings'].append(f'{type(exc).__name__}: {str(exc)[:200]}')
    out['text']='\n'.join(b['text'] for b in out['blocks'])
    if out['status']=='ok' and not out['text'].strip(): out['status']='empty'
    return out


def parse_document(data: bytes, url: str, content_type: str = '') -> dict:
    return _parse(data,url,content_type,0)
