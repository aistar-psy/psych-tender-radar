import io
import unittest
import zipfile
from radar.documents import parse_document


class DocumentTests(unittest.TestCase):
    def test_qr_sidebar_and_intent_boilerplate_are_not_full_bodies(self):
        raw='<title>学校心理设备招标公告</title><div class="content">扫码关注公众号 招标商机 手机实时看</div><div class="bid-content">下文中****为隐藏内容，仅对会员开放。项目名称：****；预算金额：50万元。</div>'
        self.assertEqual(parse_document(raw.encode(),'https://www.qianlima.com/bid-1.html')['content_status'],'unavailable')
        raw='<title>学校心理设备采购意向</title><div id="content">本次公开的采购意向是本单位政府采购工作的初步安排，具体采购项目情况以相关采购公告和采购文件为准。</div>'
        self.assertEqual(parse_document(raw.encode(),'https://example.edu.cn/intent.html')['content_status'],'unavailable')

    def test_shaanxi_news_and_company_registry_are_not_notices(self):
        for channel in ('xwzx','zcfg','cxgk'):
            doc=parse_document('<article><h1>心理服务公司</h1><p>公共资源交易信息</p></article>'.encode(),f'https://www.sxggzyjy.cn/{channel}/a.html')
            self.assertEqual(doc.get('page_kind'),'non_notice')

    def test_shaanxi_notice_container_beats_address_span(self):
        raw='<title>全国公共资源交易平台（陕西省）陕西省公共资源交易中心</title><div class="ewb-article-title">某学校心理健康设备采购公告</div><div class="epoint-article-content"><div id="noticeArea"><p>项目编号：X-26</p><p>预算金额：50万元</p><p>文件获取地址：<span class="content">某市采购中心</span></p></div></div>'
        doc=parse_document(raw.encode(),'https://www.sxggzyjy.cn/notice.html')
        self.assertIn('预算金额：50万元',doc['text'])
        self.assertEqual(doc['title'],'某学校心理健康设备采购公告')

    def test_navigation_downloads_are_not_procurement_attachments(self):
        raw='<article><h1>学校心理设备采购公告</h1><p>项目编号：X1</p><a href="/files/need.pdf">采购需求.pdf</a></article><a href="/app">APP 下载</a><a href="/help">下载中心</a>'
        doc=parse_document(raw.encode(),'https://example.org/a.html')
        self.assertEqual([x['url'] for x in doc['links'] if x['kind']=='attachment'], ['https://example.org/files/need.pdf'])

    def test_public_embedded_body_enters_attachment_followup(self):
        doc=parse_document('<article><h1>学校心理采购公告</h1><iframe src="/public/body.html"></iframe></article>'.encode(),'https://example.org/a.html')
        self.assertEqual(doc['links'][0]['kind'],'attachment')

    def test_member_summary_is_not_acquired_fulltext(self):
        raw='<article><h1>某学校心理健康设备中标公告</h1><table><tr><td>采购项目名称</td><td>心理健康设备</td></tr><tr><td>中标金额</td><td>241万元</td></tr></table><p>本网站会员请登录</p><p>查看政府采购详细信息，注册会员请点击此处</p></article>'
        doc=parse_document(raw.encode(),'https://example.org/result.html')
        self.assertEqual(doc['content_status'],'unavailable')

    def test_search_widget_cannot_replace_real_notice_body(self):
        raw='''<title>陕西省政府采购网</title><div class="content"><form>搜标题 搜全文 /</form></div><h1 class="info-title">某市学校心理中心设备采购意向</h1><div>信息时间：2026-04-01 17:05</div><div id="content"><p>采购项目名称：心理中心设备采购</p><p>采购预算：70万元</p></div>'''
        doc=parse_document(raw.encode(),'https://e/notice.html')
        self.assertIn('采购预算：70万元',doc['text'])
        self.assertNotIn('搜全文',doc['text'])
        self.assertEqual(doc['title'],'某市学校心理中心设备采购意向')
        self.assertTrue(any(b['kind']=='metadata' and '信息时间' in b['text'] for b in doc['blocks']))

    def test_empty_body_retains_title_and_public_attachment_gap(self):
        raw='''<title>河南省政府采购网</title><h1>某学院心理健康中心建设项目竞争性磋商公告</h1><div>发布时间：2026-08-12</div><div class="Content"><div id="content"></div></div><a href="/files/notice.pdf">WJ.pdf</a><div>相关新闻：校园新闻</div>'''
        doc=parse_document(raw.encode(),'https://e/notice.html')
        self.assertEqual(doc['status'],'partial')
        self.assertEqual(doc['content_status'],'unavailable')
        self.assertEqual(doc['title'],'某学院心理健康中心建设项目竞争性磋商公告')
        self.assertNotIn('相关新闻',doc['text'])
        self.assertEqual(doc['links'][0]['url'],'https://e/files/notice.pdf')

    def test_contact_heading_cannot_become_notice_title(self):
        raw='''<title>学校心理咨询室项目成交结果公告</title><h2 class="tc">学校心理咨询室项目成交结果公告</h2><div class="vF_detail_content"><h2>1.采购人信息</h2><p>名称：某学校</p><h2>3. 项目联系方式</h2><p>预算金额：20万元</p></div>'''
        self.assertEqual(parse_document(raw.encode(),'https://e/a.htm')['title'],'学校心理咨询室项目成交结果公告')

    def test_generic_title_falls_back_to_explicit_project_name(self):
        for content in ('<p>二、合同名称：学校心理中心设备采购项目</p>', '<table><tr><td>序号</td><td>采购单位名称</td><td>采购项目名称</td></tr><tr><td>1</td><td>某学院</td><td>学校心理中心设备采购项目</td></tr></table>'):
            doc=parse_document(('<html><title>公告内容文档</title><body>'+content+'</body></html>').encode(),'https://e/a.htm')
            self.assertEqual(doc['title'],'学校心理中心设备采购项目')
            self.assertTrue(doc.get('title_evidence',{}).get('locator'))

    def test_nested_procurement_table_keeps_column_boundaries(self):
        raw='''<title>公告内容文档</title><table><tr><td>本次公开采购意向</td></tr><tr><td><table><tr><td>序号</td><td>采购单位名称</td><td>采购项目名称</td><td>采购需求概况</td></tr><tr><td>1</td><td>某学院</td><td>某学院心理中心采购项目</td><td>心理设备及软件一批</td></tr></table></td></tr></table>'''
        doc=parse_document(raw.encode(),'https://e/a.htm')
        self.assertIn('1 | 某学院 | 某学院心理中心采购项目 | 心理设备及软件一批',doc['text'])
        self.assertEqual(doc['title'],'某学院心理中心采购项目')
        self.assertEqual(doc['content_status'],'available')

    def test_iframe_is_only_public_link_and_remains_a_body_gap(self):
        raw='''<article><h1>学校心理设备采购公告</h1><iframe src="/public/body.html"></iframe><iframe src="https://127.0.0.1/a"></iframe><iframe src="https://user:secret@example.org/a"></iframe><iframe src="http://localhost/a"></iframe></article>'''
        doc=parse_document(raw.encode(),'https://e/a.html')
        self.assertEqual(doc['content_status'],'unavailable')
        self.assertEqual([link['url'] for link in doc['links']],['https://e/public/body.html'])
        self.assertTrue(any('iframe' in warning for warning in doc['warnings']))

    def test_national_style_notice_excludes_qq_help_widget(self):
        raw='''<title>学校心理设备中标公告</title><div id="zbDiv"><h2 class="tc">学校心理设备中标公告</h2><table><tr><td>采购人</td><td>某市第一中学</td></tr></table><div class="vT_detail_content">中标金额：20万元</div></div><div class="fixedbox"><p>QQ服务群(工作日)</p><p>采购人:498912314</p></div>'''
        doc=parse_document(raw.encode(),'https://e/notice.jsp')
        self.assertIn('采购人 | 某市第一中学',doc['text'])
        self.assertNotIn('498912314',doc['text'])

    def test_generic_listing_is_not_available_notice_body(self):
        raw='''<html><title>招标公告</title><body><div><a href="/notice?id=1">学校心理服务采购公告</a></div><div class="fixedbox">采购人:498912314</div></body></html>'''
        doc=parse_document(raw.encode(),'https://e/index.jsp')
        self.assertEqual(doc['content_status'],'unavailable')

    def test_html_body_tables_and_links(self):
        raw = '''<html><title>公告</title><nav>首页 导航</nav><article><h1>校园心理采购</h1><p>预算金额：20万元</p><table><tr><td rowspan="2">心理测评系统</td><td>并发≥100</td></tr><tr><td>量表≥20</td></tr></table><a href="files/a.pdf">采购附件</a></article><footer>备案</footer></html>'''
        doc = parse_document(raw.encode(), 'https://example.org/a.html')
        self.assertEqual(doc['status'], 'ok')
        self.assertNotIn('导航', doc['text'])
        self.assertEqual(doc['title'], '校园心理采购')
        self.assertIn('心理测评系统 | 量表≥20', doc['text'])
        self.assertTrue(any('row:2' in b['locator'] for b in doc['blocks']))
        self.assertEqual(doc['links'][0]['url'], 'https://example.org/files/a.pdf')
        self.assertEqual(doc['links'][0]['kind'], 'attachment')

    def test_docx_and_xlsx_merged_cells(self):
        from docx import Document
        from openpyxl import Workbook
        d = Document(); d.add_paragraph('心理设备'); t = d.add_table(rows=2, cols=2)
        t.cell(0,0).text='设备'; t.cell(0,0).merge(t.cell(1,0)); t.cell(1,1).text='数量 2'
        buf=io.BytesIO(); d.save(buf)
        doc=parse_document(buf.getvalue(), 'https://e/a.docx')
        self.assertIn('数量 2', doc['text'])
        self.assertTrue(any('table:1/row:2' in b['locator'] for b in doc['blocks']))
        w=Workbook(); s=w.active; s['A1']='设备'; s.merge_cells('A1:A2'); s['B2']='规格≥3'
        buf=io.BytesIO(); w.save(buf)
        doc=parse_document(buf.getvalue(), 'https://e/a.xlsx')
        self.assertIn('设备 | 规格≥3',doc['text'])

    def test_zip_rejects_paths_and_keeps_member_locator(self):
        buf=io.BytesIO()
        with zipfile.ZipFile(buf,'w') as z:
            z.writestr('../outside.txt','bad')
            z.writestr('采购.txt','心理测评：20套')
        doc=parse_document(buf.getvalue(),'https://e/a.zip')
        self.assertEqual(doc['status'],'partial')
        self.assertNotIn('bad',doc['text'])
        self.assertTrue(any('采购.txt' in b['locator'] for b in doc['blocks']))
        self.assertTrue(doc['warnings'])

    def test_invalid_pdf_explicit_failure(self):
        doc=parse_document(b'%PDF-1.4 invalid','https://e/a.pdf')
        self.assertIn(doc['status'], ['dependency_missing','parse_error'])
        self.assertTrue(doc['warnings'])

    def test_binary_unsupported_not_decoded_as_text(self):
        self.assertEqual(parse_document(b'\x00\x01\xff','https://e/a.doc')['status'],'unsupported')

    def test_blank_pdf_is_needs_ocr(self):
        from pypdf import PdfWriter
        writer=PdfWriter(); writer.add_blank_page(width=100,height=100)
        data=io.BytesIO(); writer.write(data)
        doc=parse_document(data.getvalue(),'https://e/scan.pdf')
        self.assertEqual(doc['status'],'needs_ocr'); self.assertEqual(doc['text'],'')

    def test_zip_member_limit_is_explicit(self):
        data=io.BytesIO()
        with zipfile.ZipFile(data,'w') as z:
            for i in range(101): z.writestr(f'{i}.txt','x')
        self.assertEqual(parse_document(data.getvalue(),'https://e/many.zip')['status'],'parse_error')

    def test_procurement_article_title_beats_column_heading(self):
        raw='''<html><h2>采购公告</h2><a>学校主页</a><div class="article"><h2 class="tit">某学院智能心理服务项目竞争性磋商公告</h2><div><p>项目编号：TEST-2026</p><p>采购人信息</p><p>名 称：某学院</p><table><tr><td>序号</td><td colspan="2">标的名称</td><td>数量</td><td>单位</td></tr><tr><td>1</td><td>AI访谈评估</td><td>干预练习</td><td>1</td><td>套</td></tr></table></div></div><aside>导航心理课程系统</aside></html>'''
        doc=parse_document(raw.encode(),'https://e/detail.htm')
        self.assertEqual(doc['title'],'某学院智能心理服务项目竞争性磋商公告')
        self.assertNotIn('学校主页',doc['text']); self.assertNotIn('导航心理课程系统',doc['text'])

    def test_download_url_uses_office_mime_and_signature(self):
        from docx import Document
        d=Document(); d.add_paragraph('心理采购附件'); b=io.BytesIO(); d.save(b)
        for mime in ('application/vnd.openxmlformats-officedocument.wordprocessingml.document','application/octet-stream'):
            self.assertIn('心理采购附件',parse_document(b.getvalue(),'https://e/download?id=1',mime)['text'])

    def test_external_metadata_and_attachment_are_preserved_without_navigation(self):
        raw='''<html><nav>学校新闻导航</nav><div>发布时间：2026-09-01</div><article><h1>校园心理采购公告</h1><p>采购内容</p></article><div><a href="/a.pdf">采购附件</a></div></html>'''
        doc=parse_document(raw.encode(),'https://e/a.html')
        self.assertNotIn('学校新闻导航',doc['text'])
        self.assertTrue(any(b['kind']=='metadata' and '2026-09-01' in b['text'] for b in doc['blocks']))
        self.assertEqual(doc['links'][0]['url'],'https://e/a.pdf')

    def test_external_split_publish_metadata_keeps_value(self):
        from radar.analysis import analyze_document
        raw='<div>发布时间：<span>2026-09-06</span></div><article><h1>学校心理采购公告</h1><p>正文</p></article>'
        doc=parse_document(raw.encode(),'https://e/a.html')
        self.assertTrue(any(b['text']=='发布时间： 2026-09-06' and b['kind']=='metadata' for b in doc['blocks']))
        self.assertEqual(analyze_document(doc,'https://e/a.html','2026-09-07T09:00:00+08:00')['published_at'],'2026-09-06')

    def test_publication_time_widgets_are_metadata_not_service_time(self):
        from radar.analysis import analyze_document
        for widget in ('<p class="conttime">时间：2026-03-19 来源：学校</p>', '<div class="titx">时间：2026-03-19 浏览次数：3</div>', '<time datetime="2026-03-19">2026-03-19</time>'):
            doc=parse_document(('<article><h1>心理采购公告</h1>'+widget+'<p class="time">服务时间：2026-04-01</p></article>').encode(),'https://e/a.html')
            self.assertEqual(analyze_document(doc,'https://e/a.html','2026-09-07T09:00:00+08:00')['published_at'],'2026-03-19')
            self.assertFalse(any('服务时间' in b['text'] for b in doc['blocks'] if b['kind']=='metadata'))
